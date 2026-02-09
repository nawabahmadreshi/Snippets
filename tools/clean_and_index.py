#!/usr/bin/env python3
"""
clean_and_index.py

Zendesk HTML cleanup + stable heading IDs + cross-reference rewriting + Excel export.

What it does:
1) Cleans Zendesk/WYSIWYG noise:
   - removes <zd-html-block> wrappers (keeps content)
   - unwraps <span> wrappers (keeps text)
   - strips unwanted attributes (style, class, data-*, rel, target, etc.)
   - removes <script> and <style> tags
2) Replaces Zendesk-generated heading IDs (e.g., h_01F...) with stable slug IDs
3) Rewrites internal cross-reference links to match new IDs:
   - href="#old" -> "#new"
   - href="something#old" -> "something#new"
4) Exports headings to Excel with absolute paths
5) Adds a "Broken Links" sheet to Excel for unresolved anchors

Usage:
  python clean_and_index.py input.html output.cleaned.html headings.xlsx [page_url]

Example:
  python clean_and_index.py guide.html guide.cleaned.html headings.xlsx https://docs.example.com/guide.html
"""

import json
import re
import sys
from pathlib import Path
from bs4 import BeautifulSoup, Comment
from openpyxl import Workbook

HEADING_TAGS = ["h1", "h2", "h3", "h4", "h5", "h6"]

# Keep only these attributes (per tag). Everything else gets removed.
ALLOWED_ATTRS = {
    "*": {"id"},
    "a": {"href", "title"},
    "img": {"src", "alt", "title"},
    "table": {"border"},
    "td": {"colspan", "rowspan"},
    "th": {"colspan", "rowspan"},
}

# Remove wrapper tags but keep their contents
UNWRAP_TAGS = {"zd-html-block", "span"}

# Remove these tags entirely (and their contents)
DROP_TAGS = {"script", "style"}


def slugify(text: str) -> str:
    text = (text or "").strip().lower()
    text = re.sub(r"\u00a0", " ", text)  # nbsp
    text = re.sub(r"[^a-z0-9]+", "-", text)
    text = re.sub(r"-{2,}", "-", text).strip("-")
    return text or "section"


def unique_id(base: str, used: set[str]) -> str:
    if base not in used:
        used.add(base)
        return base
    i = 2
    while f"{base}-{i}" in used:
        i += 1
    final = f"{base}-{i}"
    used.add(final)
    return final


def heading_level(tag_name: str) -> int:
    return int(tag_name[1])


def is_zendesk_generated_id(val: str) -> bool:
    # Matches ids like h_01F2GC36KSSZJMS4VPK8SSSZ7B
    return bool(re.fullmatch(r"h_[0-9A-Z]{20,}", val or ""))


def clean_html(soup: BeautifulSoup) -> None:
    # Remove HTML comments
    for c in soup.find_all(string=lambda x: isinstance(x, Comment)):
        c.extract()

    # Drop unwanted tags
    for t in soup.find_all(DROP_TAGS):
        t.decompose()

    # Unwrap wrapper tags (keep inner content)
    for tag_name in list(UNWRAP_TAGS):
        for t in soup.find_all(tag_name):
            t.unwrap()

    # Normalize NBSP in text nodes
    for text_node in soup.find_all(string=True):
        if "\u00a0" in text_node:
            text_node.replace_with(text_node.replace("\u00a0", " "))

    # Strip attributes aggressively
    for el in soup.find_all(True):
        allowed = set(ALLOWED_ATTRS.get("*", set())) | set(ALLOWED_ATTRS.get(el.name, set()))
        for attr in list(el.attrs.keys()):
            if attr not in allowed:
                del el.attrs[attr]


def rewrite_anchor_href(href: str, id_map: dict[str, str]) -> tuple[str, bool]:
    """
    Returns (new_href, changed?)
    - Rewrites #old -> #new if old in id_map
    - Rewrites anything#old -> anything#new
    """
    if not href or "#" not in href:
        return href, False

    base, frag = href.rsplit("#", 1)
    if not frag:
        return href, False

    if frag in id_map:
        new_href = f"{base}#{id_map[frag]}" if base else f"#{id_map[frag]}"
        return new_href, True

    return href, False


def add_heading_ids_and_collect(
    soup: BeautifulSoup,
    page_url: str = "",
    replace_existing_ids: bool = True,
):
    """
    Adds stable IDs to headings and returns:
      - rows: list of dicts for Excel/JSON
      - id_map: mapping old_id -> new_id (for rewriting cross references)
    """
    used_ids = set()
    for el in soup.find_all(True):
        if el.has_attr("id") and el["id"]:
            used_ids.add(el["id"])

    stack: dict[int, tuple[str, str]] = {}
    rows: list[dict] = []
    id_map: dict[str, str] = {}

    for h in soup.find_all(HEADING_TAGS):
        title = h.get_text(" ", strip=True)

        current = (h.get("id") or "").strip()
        should_replace = replace_existing_ids and (not current or is_zendesk_generated_id(current))

        if should_replace:
            old = current
            base = slugify(title)
            hid = unique_id(base, used_ids)
            h["id"] = hid
            if old:
                id_map[old] = hid
        else:
            hid = current or unique_id(slugify(title), used_ids)
            h["id"] = hid
            used_ids.add(hid)

        lvl = heading_level(h.name)

        stack[lvl] = (title, hid)
        for deeper in range(lvl + 1, 7):
            stack.pop(deeper, None)

        path_titles = [stack[i][0] for i in range(1, lvl + 1) if i in stack]
        abs_path = " > ".join(path_titles)

        anchor = f"#{hid}"
        link = f"{page_url}{anchor}" if page_url else anchor

        rows.append({
            "level": lvl,
            "heading": title,
            "id": hid,
            "absolutePath": abs_path,
            "anchor": anchor,
            "link": link,
        })

    return rows, id_map


def rewrite_cross_references(soup: BeautifulSoup, id_map: dict[str, str]):
    """
    Rewrites <a href> anchors based on id_map.
    Returns: (changed_count, broken_links_list)
    """
    broken = []
    changed_count = 0

    current_ids = set()
    for el in soup.find_all(True):
        if el.has_attr("id") and el["id"]:
            current_ids.add(el["id"])

    for a in soup.find_all("a"):
        href = a.get("href", "")
        if not href:
            continue

        new_href, changed = rewrite_anchor_href(href, id_map)
        if changed:
            a["href"] = new_href
            changed_count += 1
            continue

        # Track internal anchors that don't exist
        if href.startswith("#"):
            target = href[1:]
            if target and target not in current_ids:
                broken.append({"href": href, "text": a.get_text(" ", strip=True)})

        # Track url#anchor that looks like Zendesk id and doesn't exist locally
        if "#" in href:
            frag = href.rsplit("#", 1)[1]
            if frag and frag not in current_ids and is_zendesk_generated_id(frag):
                broken.append({"href": href, "text": a.get_text(" ", strip=True)})

    return changed_count, broken


def write_excel_and_json(rows: list[dict], broken_links: list[dict], xlsx_out: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Headings"

    ws.append(["Level", "Heading", "ID", "Absolute Path", "Anchor", "Link"])
    for r in rows:
        ws.append([r["level"], r["heading"], r["id"], r["absolutePath"], r["anchor"], r["link"]])

    widths = [8, 48, 28, 64, 20, 64]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    ws2 = wb.create_sheet("Broken Links")
    ws2.append(["Href", "Link Text"])
    for b in broken_links:
        ws2.append([b["href"], b["text"]])

    wb.save(xlsx_out)

    json_out = Path(xlsx_out).with_suffix(".json")
    json_out.write_text(json.dumps(rows, indent=2), encoding="utf-8")


def main(html_in: str, html_out: str, xlsx_out: str, page_url: str = ""):
    soup = BeautifulSoup(Path(html_in).read_text(encoding="utf-8"), "lxml")

    # 1) Clean Zendesk HTML
    clean_html(soup)

    # 2) Add stable heading IDs + build heading rows + build old->new id map
    rows, id_map = add_heading_ids_and_collect(soup, page_url=page_url, replace_existing_ids=True)

    # 3) Rewrite cross-references
    changed_links, broken_links = rewrite_cross_references(soup, id_map)

    # 4) Write outputs
    Path(html_out).write_text(str(soup), encoding="utf-8")
    write_excel_and_json(rows, broken_links, xlsx_out)

    print(f" Clean HTML written: {html_out}")
    print(f" Excel written:      {xlsx_out}")
    print(f" JSON written:       {Path(xlsx_out).with_suffix('.json').name}")
    print(f" Headings indexed:   {len(rows)}")
    print(f" Links rewritten:    {changed_links}")
    print(f"  Broken anchors:    {len(broken_links)} (see 'Broken Links' sheet)")


if __name__ == "__main__":
    if len(sys.argv) < 4:
        print("Usage: python clean_and_index.py input.html output.cleaned.html headings.xlsx [page_url]")
        sys.exit(1)

    html_in, html_out, xlsx_out = sys.argv[1], sys.argv[2], sys.argv[3]
    page_url = sys.argv[4] if len(sys.argv) >= 5 else ""
    main(html_in, html_out, xlsx_out, page_url)
